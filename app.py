# app.py
import csv
import os
import sqlite3
import datetime
import io
import openpyxl
import shutil
import time
import threading
from openpyxl.styles import Font
from functools import wraps
from flask import Flask, flash, g, redirect, render_template, request, url_for, jsonify, session, send_file
from flask_wtf import FlaskForm, CSRFProtect
from flask_wtf.csrf import CSRFError, generate_csrf
from wtforms import StringField, PasswordField, SelectField, TextAreaField
from wtforms.validators import DataRequired, Length, Optional
from werkzeug.security import generate_password_hash, check_password_hash

app = Flask(__name__)
app.config['DATABASE'] = 'laboratorio.db'
app.secret_key = os.environ.get('SECRET_KEY', 'dev-laboratorio-secret-key')
app.config.update(
    SESSION_COOKIE_SECURE=False,
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE='Lax',
    WTF_CSRF_ENABLED=True,
    WTF_CSRF_TIME_LIMIT=None,
    WTF_CSRF_HEADERS=['X-CSRFToken', 'X-CSRF-Token']
)
csrf = CSRFProtect(app)

@app.errorhandler(CSRFError)
def handle_csrf_error(e):
    if request.path.startswith('/api/'):
        return jsonify({'error': 'Error de seguridad (CSRF). Por favor, recarga la página.'}), 400
    return render_template('error.html', message='Error de seguridad (CSRF). Por favor, recarga la página.', code=400), 400

login_attempts = {}
login_lock = threading.Lock()
undo_delete_cache = {}

CATALOGS = {
    'estados': {'table': 'EstadosCatalogo', 'column': 'estado', 'default': 'PEND. DE REVISION', 'title': 'Estados'},
    'equipos': {'table': 'EquiposCatalogo', 'column': 'equipo', 'default': 'DESCONOCIDO', 'title': 'Equipos'},
    'salas': {'table': 'SalasCatalogo', 'column': 'sala', 'default': 'DESCONOCIDA', 'title': 'Salas'},
    'familias': {'table': 'FamiliasCatalogo', 'column': 'familia', 'default': '', 'title': 'Familias'},
    'tecnicos': {'table': 'Tecnicos', 'column': None, 'default': None, 'title': 'Técnicos'}
}

INDEX_FILTERS = [
    ('URGENTES', 'Urgentes'), ('PEND_REPARACION', 'Pend. reparación'),
    ('EN_REPARACION', 'En reparación'), ('REPARADOS', 'Reparados'),
    ('PENDIENTES', 'Pendientes'), ('SIN_REPARACION', 'Sin reparación'), ('TODAS', 'Todas')
]

# =====================================================================
# BASE DE DATOS Y TABLAS
# =====================================================================
def ensure_support_tables(db):
    db.executescript('''
        CREATE TABLE IF NOT EXISTS Reparaciones (
            id INTEGER PRIMARY KEY AUTOINCREMENT, fecha TEXT NOT NULL, sala TEXT NOT NULL,
            uid TEXT, npu TEXT, parte TEXT, familia TEXT, equipo TEXT NOT NULL,
            urgente TEXT CHECK(urgente IN ('SI', 'NO')) DEFAULT 'NO', tecnico_id INTEGER,
            estado TEXT DEFAULT 'PEND. DE REVISION', observaciones TEXT,
            dia_semana INTEGER, fecha_en_reparacion TEXT DEFAULT '', fecha_reparado TEXT DEFAULT '',
            fecha_pendiente TEXT DEFAULT '', fecha_sin_reparacion TEXT DEFAULT '',
            FOREIGN KEY (tecnico_id) REFERENCES Tecnicos (id)
        );
        CREATE TABLE IF NOT EXISTS Tecnicos (id INTEGER PRIMARY KEY AUTOINCREMENT, nombre TEXT UNIQUE NOT NULL, activo INTEGER NOT NULL DEFAULT 1);
        CREATE TABLE IF NOT EXISTS EstadosCatalogo (id INTEGER PRIMARY KEY AUTOINCREMENT, nombre TEXT UNIQUE NOT NULL);
        CREATE TABLE IF NOT EXISTS EquiposCatalogo (id INTEGER PRIMARY KEY AUTOINCREMENT, nombre TEXT UNIQUE NOT NULL, valor TEXT DEFAULT '');
        CREATE TABLE IF NOT EXISTS SalasCatalogo (id INTEGER PRIMARY KEY AUTOINCREMENT, nombre TEXT UNIQUE NOT NULL);
        CREATE TABLE IF NOT EXISTS FamiliasCatalogo (id INTEGER PRIMARY KEY AUTOINCREMENT, nombre TEXT UNIQUE NOT NULL);

        CREATE TABLE IF NOT EXISTS Usuarios (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT UNIQUE NOT NULL,
            password_hash TEXT NOT NULL,
            rol TEXT CHECK(rol IN ('admin', 'tecnico')) NOT NULL,
            tecnico_id INTEGER,
            session_id TEXT,
            last_ip TEXT,
            FOREIGN KEY (tecnico_id) REFERENCES Tecnicos (id)
        );

        CREATE TABLE IF NOT EXISTS Historial (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            usuario_id INTEGER,
            username TEXT NOT NULL,
            accion TEXT NOT NULL,
            entidad TEXT,
            entidad_id INTEGER,
            detalle TEXT,
            ip TEXT,
            fecha TEXT DEFAULT (datetime('now', 'localtime')),
            FOREIGN KEY (usuario_id) REFERENCES Usuarios(id)
        );
    ''')

def init_db():
    with app.app_context():
        db = get_db()
        ensure_support_tables(db)

        # Crear Admin por defecto si no existe
        admin = db.execute("SELECT * FROM Usuarios WHERE username = 'admin'").fetchone()
        if not admin:
            db.execute("INSERT INTO Usuarios (username, password_hash, rol) VALUES (?, ?, ?)",
                       ('admin', generate_password_hash('Admin@Lab2024!'), 'admin'))
        db.commit()

def get_db():
    if 'db' not in g:
        g.db = sqlite3.connect(app.config['DATABASE'])
        g.db.row_factory = sqlite3.Row
        
        # Función personalizada para ordenar fechas inconsistentes (D/M/YYYY o YYYY/MM/DD -> YYYYMMDD)
        def sql_sortable_date(fecha_str):
            if not fecha_str: return "00000000000000"
            try:
                f_limpia = str(fecha_str).strip().replace('-', '/').replace('.', '/')
                partes = [p for p in f_limpia.split(' ') if p.strip()]
                if not partes: return "00000000000000"
                f_parts = partes[0].split('/')
                if len(f_parts) < 3: return "00000000000000"
                
                # Detectar formato: YYYY/MM/DD vs DD/MM/YYYY
                if len(f_parts[0]) == 4:
                    y, m, d = int(f_parts[0]), int(f_parts[1]), int(f_parts[2])
                else:
                    d, m = int(f_parts[0]), int(f_parts[1])
                    y_str = f_parts[2]
                    y = int(y_str) if len(y_str) == 4 else int(y_str) + 2000
                
                h = minu = s = 0
                if len(partes) > 1:
                    t_parts = partes[1].split(':')
                    if len(t_parts) > 0 and t_parts[0].isdigit(): h = int(t_parts[0])
                    if len(t_parts) > 1 and t_parts[1].isdigit(): minu = int(t_parts[1])
                    if len(t_parts) > 2 and t_parts[2].isdigit(): s = int(t_parts[2])
                
                return f"{y:04d}{m:02d}{d:02d}{h:02d}{minu:02d}{s:02d}"
            except:
                return "00000000000000"
        
        g.db.create_function("sortable_date", 1, sql_sortable_date)
        
        # MEJORAS DE ESTABILIDAD Y SEGURIDAD (Puntos 1 y 2)
        g.db.execute('PRAGMA journal_mode = WAL;')
        g.db.execute('PRAGMA foreign_keys = ON;')
        
        ensure_support_tables(g.db)
        g.db.commit()
    return g.db

@app.teardown_appcontext
def close_connection(exception):
    db = g.pop('db', None)
    if db is not None: db.close()

# =====================================================================
# SEGURIDAD Y LOGIN
# =====================================================================
import uuid

@app.before_request
def require_login():
    rutas_libres = ['login', 'static', 'health']
    if request.endpoint not in rutas_libres and 'user_id' not in session:
        if request.path.startswith('/api/'):
            return jsonify({'error': 'Sesión expirada. Por favor, recarga la página e inicia sesión nuevamente.'}), 401
        return redirect(url_for('login'))
    
    # Verificar si la sesión es válida (única sesión)
    if 'user_id' in session and request.endpoint != 'logout':
        db = get_db()
        user = db.execute('SELECT session_id FROM Usuarios WHERE id = ?', (session['user_id'],)).fetchone()
        if user and user['session_id'] != session.get('sid'):
            session.clear()
            if request.path.startswith('/api/'):
                return jsonify({'error': 'Sesión cerrada desde otro dispositivo.'}), 401
            return redirect(url_for('login', session_closed=1))

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        ip = request.remote_addr

        allowed, msg = check_rate_limit(ip)
        if not allowed:
            flash(msg, 'danger')
            return render_template('login.html')

        username = request.form.get('username', '').strip()
        password = request.form.get('password', '')
        force = request.form.get('force') == 'true'

        if not username or not password:
            flash('Usuario y contraseña son obligatorios.', 'danger')
            return render_template('login.html')

        db = get_db()

        user = db.execute('''
            SELECT u.*, t.nombre as tecnico_nombre
            FROM Usuarios u
            LEFT JOIN Tecnicos t ON u.tecnico_id = t.id
            WHERE u.username = ?
        ''', (username,)).fetchone()

        if user and check_password_hash(user['password_hash'], password):
            reset_rate_limit(ip)

            current_ip = request.remote_addr
            if user['session_id'] and not force:
                return render_template('login.html',
                                     error_session=True,
                                     ip_conectada=user['last_ip'],
                                     username=username,
                                     password=password)

            sid = str(uuid.uuid4())
            db.execute('UPDATE Usuarios SET session_id = ?, last_ip = ? WHERE id = ?',
                       (sid, current_ip, user['id']))
            db.commit()

            session.clear()
            session['user_id'] = user['id']
            session['username'] = user['username']
            session['rol'] = user['rol']
            session['tecnico_id'] = user['tecnico_id']
            session['sid'] = sid
            session['tecnico_nombre'] = user['tecnico_nombre'] if user['tecnico_nombre'] else user['username']

            registrar_historial(user['id'], user['username'], 'LOGIN', 'Sistema', None, 'Inicio de sesión exitoso')
            return redirect(url_for('index'))

        with login_lock:
            if ip in login_attempts:
                login_attempts[ip]['count'] = login_attempts[ip].get('count', 0) + 1

        flash('Usuario o contraseña incorrectos.', 'danger')
    return render_template('login.html')

@app.route('/logout')
def logout():
    if 'user_id' in session:
        user_id = session.get('user_id')
        username = session.get('username', 'unknown')
        db = get_db()
        db.execute('UPDATE Usuarios SET session_id = NULL WHERE id = ?', (user_id,))
        db.commit()
        registrar_historial(user_id, username, 'LOGOUT', 'Sistema', None, 'Cierre de sesión')
    session.clear()
    return redirect(url_for('login'))

def admin_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if session.get('rol') != 'admin':
            if request.path.startswith('/api/'):
                return jsonify({'error': 'Acceso denegado. Se requieren permisos de administrador.'}), 403
            flash('Acceso denegado. Solo administradores.', 'danger')
            return redirect(url_for('index'))
        return f(*args, **kwargs)
    return decorated_function

# =====================================================================
# FUNCIONES AUXILIARES
# =====================================================================
def ensure_catalog_value(db, table, value, valor=''):
    nombre = (value or '').strip()
    if nombre:
        if table == 'EquiposCatalogo':
            db.execute(f'INSERT OR IGNORE INTO {table} (nombre, valor) VALUES (?, ?)', (nombre, normalize_price_input(valor)))
        elif table == 'Tecnicos':
            db.execute(f'INSERT OR IGNORE INTO {table} (nombre) VALUES (?)', (nombre,))
            db.execute(f'UPDATE {table} SET activo = 1 WHERE nombre = ?', (nombre,))
        else:
            db.execute(f'INSERT OR IGNORE INTO {table} (nombre) VALUES (?)', (nombre,))

def stored_price_to_amount(valor):
    raw = str(valor or '').strip().replace('$', '').replace(' ', '')
    if not raw:
        return 0.0

    if raw.isdigit():
        return int(raw) / 100.0

    if ',' in raw:
        return float(raw.replace('.', '').replace(',', '.'))

    if raw.count('.') == 1:
        integer_part, decimal_part = raw.split('.')
        if decimal_part.isdigit() and len(decimal_part) == 2:
            return float(f"{integer_part}.{decimal_part}")

    return float(raw.replace('.', ''))

def normalize_price_input(valor):
    raw = str(valor or '').strip().replace('$', '').replace(' ', '')
    if not raw:
        return ''

    if ',' in raw:
        amount = float(raw.replace('.', '').replace(',', '.'))
    elif raw.count('.') == 1:
        integer_part, decimal_part = raw.split('.')
        if decimal_part.isdigit() and len(decimal_part) == 2:
            amount = float(f"{integer_part}.{decimal_part}")
        else:
            amount = float(raw.replace('.', ''))
    elif raw.count('.') > 1:
        amount = float(raw.replace('.', ''))
    else:
        amount = float(raw)

    cents = int(round(amount * 100))
    return str(cents)

def get_catalog_items(db, tipo):
    table = CATALOGS[tipo]['table']
    if tipo == 'equipos': 
        return db.execute(f'SELECT id, nombre, valor FROM {table} ORDER BY nombre').fetchall()
    if tipo == 'tecnicos': 
        return db.execute('''
            SELECT 
                t.id, 
                t.nombre, 
                COALESCE(t.activo, 1) as activo, 
                SUM(CASE WHEN r.id IS NOT NULL AND UPPER(COALESCE(r.estado, '')) NOT LIKE 'REPARADO%' AND UPPER(COALESCE(r.estado, '')) NOT LIKE 'SIN REPARACION%' THEN 1 ELSE 0 END) AS rep_en_su_poder,
                SUM(CASE WHEN r.id IS NOT NULL AND UPPER(COALESCE(r.estado, '')) LIKE 'REPARADO%' THEN 1 ELSE 0 END) AS reparados
            FROM Tecnicos t 
            LEFT JOIN Reparaciones r ON r.tecnico_id = t.id 
            GROUP BY t.id 
            ORDER BY t.activo DESC, t.nombre ASC
        ''').fetchall()
    return db.execute(f'SELECT id, nombre FROM {table} ORDER BY nombre').fetchall()

def calcular_dia_semana(fecha_str):
    if not fecha_str: return None
    try:
        f_limpia = fecha_str.strip().split(' ')[0] 
        partes = f_limpia.split('/') if '/' in f_limpia else f_limpia.split('-')
        y, m, d = partes if len(partes[0]) == 4 else (partes[2], partes[1], partes[0])
        dt = datetime.date(int(y), int(m), int(d))
        return dt.isoweekday() % 7 
    except: return None

@app.context_processor
def utility_processor():
    def format_price(amount):
        if not amount: return "$ 0"
        try:
            value = stored_price_to_amount(amount)
            integer_part, decimal_part = f"{value:,.2f}".split('.')
            return f"$ {integer_part.replace(',', '.')},{decimal_part}"
        except:
            return f"$ {amount}"

    def get_session():
        return dict(session)

    return dict(format_price=format_price, csrf_token=generate_csrf, current_user=get_session())

def registrar_historial(usuario_id, username, accion, entidad='', entidad_id=None, detalle=''):
    try:
        db = get_db()
        ip = request.remote_addr if request else 'unknown'
        db.execute('''INSERT INTO Historial (usuario_id, username, accion, entidad, entidad_id, detalle, ip)
                      VALUES (?, ?, ?, ?, ?, ?, ?)''',
                   (usuario_id, username, accion, entidad, entidad_id, detalle, ip))
        db.commit()
    except Exception as e:
        print(f"Error registrando historial: {e}")

def check_rate_limit(ip):
    now = time.time()
    with login_lock:
        if ip not in login_attempts:
            login_attempts[ip] = {'count': 0, 'first_attempt': now, 'blocked_until': 0}

        attempt = login_attempts[ip]

        if attempt.get('blocked_until', 0) > now:
            return False, f"Demasiados intentos. Intenta en {int(attempt['blocked_until'] - now)} segundos."

        if now - attempt['first_attempt'] > 300:
            attempt['count'] = 0
            attempt['first_attempt'] = now

        attempt['count'] = attempt.get('count', 0) + 1

        if attempt['count'] > 5:
            attempt['blocked_until'] = now + 900
            return False, "Demasiados intentos. Bloqueado por 15 minutos."

        return True, None

def reset_rate_limit(ip):
    with login_lock:
        if ip in login_attempts:
            login_attempts[ip] = {'count': 0, 'first_attempt': time.time(), 'blocked_until': 0}

def create_backup_file():
    os.makedirs('backups', exist_ok=True)
    timestamp = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
    backup_path = f'backups/laboratorio_{timestamp}.db'
    shutil.copy2('laboratorio.db', backup_path)
    return backup_path

def decode_csv_bytes(file_bytes):
    for encoding in ('utf-8-sig', 'utf-8', 'cp1252', 'latin-1'):
        try:
            return file_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError('No se pudo leer el archivo CSV con una codificación soportada.')

# Función para normalizar fechas a formato DD/MM/YYYY HH:MM
def normalize_date(fecha_str):
    if not fecha_str: return ''
    try:
        f_limpia = str(fecha_str).strip().replace('-', '/').replace('.', '/')
        partes = [p for p in f_limpia.split(' ') if p.strip()]
        if not partes: return ''
        
        f_parts = partes[0].split('/')
        if len(f_parts) < 3: return fecha_str # No se pudo parsear, devolver original
        
        # Detectar formato: YYYY/MM/DD vs DD/MM/YYYY
        if len(f_parts[0]) == 4:
            y, m, d = int(f_parts[0]), int(f_parts[1]), int(f_parts[2])
        else:
            d, m = int(f_parts[0]), int(f_parts[1])
            y_str = f_parts[2]
            y = int(y_str) if len(y_str) == 4 else int(y_str) + 2000
            
        h = minu = s = 0
        has_time = False
        if len(partes) > 1:
            t_parts = partes[1].split(':')
            if len(t_parts) > 0 and t_parts[0].isdigit(): h = int(t_parts[0]); has_time = True
            if len(t_parts) > 1 and t_parts[1].isdigit(): minu = int(t_parts[1])
            if len(t_parts) > 2 and t_parts[2].isdigit(): s = int(t_parts[2])
            
        if has_time:
            return f"{d:02d}/{m:02d}/{y:04d} {h:02d}:{minu:02d}"
        return f"{d:02d}/{m:02d}/{y:04d}"
    except:
        return fecha_str

# Función auxiliar para obtener valores del CSV con manejo de nombres de columnas alternativos
def get_csv_value(row, key, default=''):
    aliases = {
        'estados de reparacion': ['estados de reparacion', 'estado', 'estados'],
        'observaciones': ['observaciones', 'obs.', 'obs'],
        'f. pendiente': ['f. pendiente', 'fecha "pendiente"', 'fecha pendiente'],
        'f. en rep.': ['f. en rep.', 'fecha "en reparacion / en prueba "', 'fecha en reparacion'],
        'f. reparado': ['f. reparado', 'fecha "reparado"', 'fecha reparado'],
        'f. sin rep.': ['f. sin rep.', 'fecha "sin reparacion"', 'fecha sin reparacion']
    }
    
    # Primero intentar con la clave exacta
    if key in row:
        val = row[key]
        return str(val).strip() if val is not None else default
        
    # Luego intentar con alias si existen
    if key in aliases:
        for alias in aliases[key]:
            if alias in row:
                val = row[alias]
                return str(val).strip() if val is not None else default
                
    return default

# =====================================================================
# RUTAS PRINCIPALES (REPARACIONES)
# =====================================================================
@app.route('/')
def index():
    db = get_db()
    
    tabs_estado = list(INDEX_FILTERS)
    is_tecnico = session.get('rol') == 'tecnico'
    if is_tecnico:
        tabs_estado.insert(1, ('MIS_REPARACIONES', 'Mis rep.'))
        
    estados_validos = {estado for estado, _ in tabs_estado}
    
    estado_defecto = 'MIS_REPARACIONES' if is_tecnico else 'PEND_REPARACION'
    estado_actual = request.args.get('estado', estado_defecto, type=str).upper()
    if estado_actual not in estados_validos: 
        estado_actual = estado_defecto

    busqueda = request.args.get('q', '', type=str).strip()
    f_sala, f_tecnico, f_estado = request.args.get('f_sala', ''), request.args.get('f_tecnico', ''), request.args.get('f_estado', '')
    f_anio, f_mes, f_dia, f_diasemana = request.args.get('f_anio', ''), request.args.get('f_mes', ''), request.args.get('f_dia', ''), request.args.get('f_diasemana', '')
    
    page = max(1, request.args.get('page', 1, type=int))
    per_page = 100

    base_where_clauses, base_query_params = [], []
    if f_sala: base_where_clauses.append('r.sala = ?'); base_query_params.append(f_sala)
    if f_tecnico:
        if f_tecnico == 'SIN_ASIGNAR': base_where_clauses.append('r.tecnico_id IS NULL')
        else: base_where_clauses.append('t.nombre = ?'); base_query_params.append(f_tecnico)
    if f_estado: base_where_clauses.append('r.estado = ?'); base_query_params.append(f_estado)
    if f_anio: base_where_clauses.append('r.fecha LIKE ?'); base_query_params.append(f'%/{f_anio}%')
    if f_mes: base_where_clauses.append('(r.fecha LIKE ? OR r.fecha LIKE ?)'); base_query_params.extend([f'%/{int(f_mes):02d}/%', f'%/{int(f_mes)}/%'])
    if f_dia: base_where_clauses.append('(r.fecha LIKE ? OR r.fecha LIKE ?)'); base_query_params.extend([f'{int(f_dia):02d}/%', f'{int(f_dia)}/%'])
    if f_diasemana: base_where_clauses.append('r.dia_semana = ?'); base_query_params.append(int(f_diasemana))
    if busqueda:
        base_where_clauses.append('(' + ' OR '.join(['r.fecha LIKE ?', 'r.sala LIKE ?', 'r.uid LIKE ?', 'r.npu LIKE ?', 'r.equipo LIKE ?', 'r.estado LIKE ?', 'r.observaciones LIKE ?', 't.nombre LIKE ?']) + ')')
        base_query_params.extend([f'%{busqueda}%'] * 8)

    base_where_sql = (' WHERE ' + ' AND '.join(base_where_clauses)) if base_where_clauses else ''

    tecnico_id_actual = session.get('tecnico_id') or 0

    counts_query = f'''
        SELECT
            SUM(CASE WHEN r.urgente = 'SI' AND UPPER(r.estado) NOT LIKE 'REPARADO%' AND UPPER(r.estado) NOT LIKE 'SIN REPARACION%' THEN 1 ELSE 0 END) as URGENTES,
            SUM(CASE WHEN UPPER(r.estado) = 'PEND. DE REVISION' THEN 1 ELSE 0 END) as PEND_REPARACION,
            SUM(CASE WHEN UPPER(r.estado) = 'EN REPARACION' THEN 1 ELSE 0 END) as EN_REPARACION,
            SUM(CASE WHEN UPPER(r.estado) LIKE 'REPARADO%' THEN 1 ELSE 0 END) as REPARADOS,
            SUM(CASE WHEN UPPER(r.estado) NOT LIKE 'REPARADO%' AND UPPER(r.estado) NOT LIKE 'SIN REPARACION%' AND UPPER(r.estado) != 'PEND. DE REVISION' AND UPPER(r.estado) != 'EN REPARACION' THEN 1 ELSE 0 END) as PENDIENTES,
            SUM(CASE WHEN UPPER(r.estado) LIKE 'SIN REPARACION%' THEN 1 ELSE 0 END) as SIN_REPARACION,
            COUNT(r.id) as TODAS,
            SUM(CASE WHEN r.tecnico_id = {tecnico_id_actual} AND UPPER(r.estado) NOT LIKE 'REPARADO%' AND UPPER(r.estado) NOT LIKE 'SIN REPARACION%' THEN 1 ELSE 0 END) as MIS_REPARACIONES
        FROM Reparaciones r LEFT JOIN Tecnicos t ON r.tecnico_id = t.id {base_where_sql}
    '''
    raw_counts = db.execute(counts_query, base_query_params).fetchone()
    tab_counts = {k: (raw_counts[k] or 0) for k in ['URGENTES', 'PEND_REPARACION', 'EN_REPARACION', 'REPARADOS', 'PENDIENTES', 'SIN_REPARACION', 'TODAS', 'MIS_REPARACIONES']}

    final_where_clauses, final_query_params = list(base_where_clauses), list(base_query_params)
    
    if estado_actual == 'URGENTES': 
        final_where_clauses.extend(["r.urgente = 'SI'", "UPPER(r.estado) NOT LIKE 'REPARADO%'", "UPPER(r.estado) NOT LIKE 'SIN REPARACION%'"])
    elif estado_actual == 'MIS_REPARACIONES': 
        final_where_clauses.extend(["r.tecnico_id = ?", "UPPER(r.estado) NOT LIKE 'REPARADO%'", "UPPER(r.estado) NOT LIKE 'SIN REPARACION%'"])
        final_query_params.append(tecnico_id_actual)
    elif estado_actual == 'PEND_REPARACION': 
        final_where_clauses.append('UPPER(r.estado) = ?'); final_query_params.append('PEND. DE REVISION')
    elif estado_actual == 'EN_REPARACION': 
        final_where_clauses.append('UPPER(r.estado) = ?'); final_query_params.append('EN REPARACION')
    elif estado_actual == 'REPARADOS': 
        final_where_clauses.append('UPPER(r.estado) LIKE ?'); final_query_params.append('REPARADO%')
    elif estado_actual == 'SIN_REPARACION': 
        final_where_clauses.append('UPPER(r.estado) LIKE ?'); final_query_params.append('SIN REPARACION%')
    elif estado_actual == 'PENDIENTES': 
        final_where_clauses.extend(["UPPER(r.estado) NOT LIKE 'REPARADO%'", "UPPER(r.estado) NOT LIKE 'SIN REPARACION%'", "UPPER(r.estado) != 'PEND. DE REVISION'", "UPPER(r.estado) != 'EN REPARACION'"])

    # Determinar qué columna de fecha usar para el ordenamiento
    order_col = "r.fecha_reparado" if estado_actual == 'REPARADOS' else "r.fecha"
    
    final_where_sql = (' WHERE ' + ' AND '.join(final_where_clauses)) if final_where_clauses else ''
    total_reparaciones = db.execute(f'SELECT COUNT(*) AS total FROM Reparaciones r LEFT JOIN Tecnicos t ON r.tecnico_id = t.id {final_where_sql}', final_query_params).fetchone()['total']
    total_pages = max(1, (total_reparaciones + per_page - 1) // per_page)
    if page > total_pages: page = total_pages
    offset = (page - 1) * per_page

    # EJECUCIÓN OPTIMIZADA: Usando la función personalizada sortable_date para un ordenamiento perfecto
    query_reparaciones = f'''
        SELECT r.*, t.nombre AS tecnico_nombre, ec.valor AS ahorro_valor
        FROM Reparaciones r 
        LEFT JOIN Tecnicos t ON r.tecnico_id = t.id 
        LEFT JOIN EquiposCatalogo ec ON ec.nombre = r.equipo
        {final_where_sql} 
        ORDER BY sortable_date({order_col}) DESC, r.id DESC
        LIMIT ? OFFSET ?
    '''
    params_final_with_limit = final_query_params + [per_page, offset]
    reparaciones_rows = db.execute(query_reparaciones, params_final_with_limit).fetchall()
    reparaciones = [dict(row) for row in reparaciones_rows]

    salas_opt = db.execute('SELECT DISTINCT sala FROM Reparaciones WHERE sala IS NOT NULL AND sala != "" ORDER BY sala').fetchall()
    tecnicos_opt = db.execute('SELECT DISTINCT t.nombre FROM Reparaciones r JOIN Tecnicos t ON r.tecnico_id = t.id ORDER BY t.nombre').fetchall()
    estados_opt = db.execute('SELECT DISTINCT estado FROM Reparaciones WHERE estado IS NOT NULL AND estado != "" ORDER BY estado').fetchall()
    
    fechas_raw = db.execute('SELECT DISTINCT fecha FROM Reparaciones WHERE fecha IS NOT NULL AND fecha != ""').fetchall()
    anios_set = set()
    for row in fechas_raw:
        try:
            f_str = row['fecha'].strip().split(' ')[0]
            partes = f_str.split('/') if '/' in f_str else f_str.split('-')
            anio = partes[0] if len(partes[0]) == 4 else (partes[2] if len(partes)==3 else None)
            if anio and len(anio) == 4 and anio.isdigit(): anios_set.add(anio)
        except: pass

    # --- LÓGICA DINÁMICA: NPU Repetidos ---
    repetidos_raw = db.execute('''
        SELECT npu FROM Reparaciones 
        WHERE npu IS NOT NULL AND TRIM(npu) != '' 
        GROUP BY TRIM(npu) 
        HAVING COUNT(id) > 1
    ''').fetchall()
    npus_repetidos = {row['npu'].strip() for row in repetidos_raw if row['npu']}

    return render_template('index.html', reparaciones=reparaciones, page=page, total_pages=total_pages, estado_actual=estado_actual, tabs_estado=tabs_estado, busqueda=busqueda, f_sala=f_sala, f_tecnico=f_tecnico, f_estado=f_estado, f_anio=f_anio, f_mes=f_mes, f_dia=f_dia, f_diasemana=f_diasemana, tab_counts=tab_counts, salas_opt=salas_opt, tecnicos_opt=tecnicos_opt, estados_opt=estados_opt, anios_opt=[{'anio': a} for a in sorted(list(anios_set), reverse=True)], npus_repetidos=npus_repetidos)

@app.route('/reparacion/nueva', methods=['GET', 'POST'])
def nueva_reparacion():
    db = get_db()
    if request.method == 'POST':
        fecha = request.form.get('fecha')
        sala = request.form.get('sala')
        uid = request.form.get('uid')
        npu = request.form.get('npu')
        familia = request.form.get('familia')
        equipo = request.form.get('equipo')
        urgente = request.form.get('urgente', 'NO')
        tecnico_id = request.form.get('tecnico_id') or None
        observaciones = request.form.get('observaciones', '')

        if session.get('rol') == 'tecnico':
            fecha = datetime.datetime.now().strftime('%d/%m/%Y %H:%M')
        
        dia_semana = calcular_dia_semana(fecha)
        estado_inicial = 'EN REPARACION' if tecnico_id else 'PEND. DE REVISION'
        timestamp = datetime.datetime.now().strftime('%d/%m/%Y %H:%M')
        
        if estado_inicial == 'PEND. DE REVISION':
            fecha_pendiente = timestamp
            fecha_en_reparacion = ''
        else:
            fecha_pendiente = timestamp
            fecha_en_reparacion = timestamp

        db.execute('''
            INSERT INTO Reparaciones 
            (fecha, sala, uid, npu, familia, equipo, urgente, tecnico_id, estado, observaciones, dia_semana, fecha_pendiente, fecha_en_reparacion)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', (fecha, sala, uid, npu, familia, equipo, urgente, tecnico_id, estado_inicial, observaciones, dia_semana, fecha_pendiente, fecha_en_reparacion))

        ensure_catalog_value(db, 'SalasCatalogo', sala)
        ensure_catalog_value(db, 'FamiliasCatalogo', familia)
        ensure_catalog_value(db, 'EquiposCatalogo', equipo)

        db.commit()
        rep_id = db.execute('SELECT last_insert_rowid()').fetchone()[0]
        registrar_historial(session['user_id'], session['username'], 'CREAR', 'Reparacion', rep_id, f'Nueva reparación: {equipo} en {sala}')
        flash('Reparación cargada con éxito.', 'success')
        return redirect(url_for('index'))

    salas = get_catalog_items(db, 'salas')
    familias = get_catalog_items(db, 'familias')
    equipos = get_catalog_items(db, 'equipos')
    tecnicos = db.execute('SELECT id, nombre FROM Tecnicos WHERE activo = 1 ORDER BY nombre').fetchall()
    hoy = datetime.datetime.now().strftime('%d/%m/%Y %H:%M')
    
    return render_template('nueva_reparacion.html', salas=salas, familias=familias, equipos=equipos, tecnicos=tecnicos, hoy=hoy)

@app.route('/reparacion/<int:id>')
def detalle_reparacion(id):
    db = get_db()
    r = db.execute('''
        SELECT r.*, t.nombre as tecnico_nombre 
        FROM Reparaciones r 
        LEFT JOIN Tecnicos t ON r.tecnico_id = t.id 
        WHERE r.id = ?
    ''', (id,)).fetchone()
    
    if not r:
        flash('Reparación no encontrada.', 'danger')
        return redirect(url_for('index'))
        
    tecnicos = db.execute('SELECT id, nombre FROM Tecnicos WHERE activo = 1 ORDER BY nombre').fetchall()
    estados = db.execute('SELECT nombre FROM EstadosCatalogo ORDER BY nombre').fetchall()
    
    salas = get_catalog_items(db, 'salas')
    familias = get_catalog_items(db, 'familias')
    equipos = get_catalog_items(db, 'equipos')

    historial_npu = []
    if r['npu'] and str(r['npu']).strip():
        npu_limpio = str(r['npu']).strip()
        historial_npu = db.execute('''
            SELECT r.id, r.fecha, r.estado, r.urgente, r.observaciones, t.nombre as tecnico_nombre 
            FROM Reparaciones r
            LEFT JOIN Tecnicos t ON r.tecnico_id = t.id
            WHERE r.npu = ? AND r.id != ? 
            ORDER BY r.id DESC
        ''', (npu_limpio, id)).fetchall()
    
    return render_template('detalle_reparacion.html', r=r, tecnicos=tecnicos, estados=estados, salas=salas, familias=familias, equipos=equipos, historial_npu=historial_npu)
    
@app.route('/reparacion/editar/<int:id>', methods=['GET', 'POST'])
@admin_required
def editar_reparacion(id):
    try:
        db = get_db()
        r = db.execute('SELECT * FROM Reparaciones WHERE id = ?', (id,)).fetchone()

        if not r:
            flash('Reparación no encontrada.', 'danger')
            return redirect(url_for('index'))

        if request.method == 'POST':
            fecha = request.form.get('fecha')
            sala = request.form.get('sala')
            uid = request.form.get('uid')
            npu = request.form.get('npu')
            familia = request.form.get('familia')
            equipo = request.form.get('equipo')
            urgente = request.form.get('urgente', 'NO')
            tecnico_id = request.form.get('tecnico_id') or None
            observaciones = request.form.get('observaciones', '')

            if session.get('rol') != 'admin':
                fecha = r['fecha']

            db.execute('''
                UPDATE Reparaciones
                SET fecha = ?, sala = ?, uid = ?, npu = ?, familia = ?, equipo = ?,
                    urgente = ?, tecnico_id = ?, observaciones = ?
                WHERE id = ?
            ''', (fecha, sala, uid, npu, familia, equipo, urgente, tecnico_id, observaciones, id))

            ensure_catalog_value(db, 'SalasCatalogo', sala)
            ensure_catalog_value(db, 'FamiliasCatalogo', familia)
            ensure_catalog_value(db, 'EquiposCatalogo', equipo)

            db.commit()
            registrar_historial(session['user_id'], session['username'], 'EDITAR', 'Reparacion', id, f'Editar: {equipo}')
            flash('Reparación actualizada correctamente.', 'success')
            return redirect(url_for('detalle_reparacion', id=id))

        salas = get_catalog_items(db, 'salas')
        familias = get_catalog_items(db, 'familias')
        equipos = get_catalog_items(db, 'equipos')
        tecnicos = db.execute('SELECT id, nombre FROM Tecnicos WHERE activo = 1 ORDER BY nombre').fetchall()

        return render_template('editar_reparaciones.html', r=r, salas=salas, familias=familias, equipos=equipos, tecnicos=tecnicos)
    except Exception as e:
        flash(f'Error: {str(e)}', 'danger')
        return redirect(url_for('index'))

@app.route('/reparacion/eliminar/<int:id>', methods=['POST'])
@admin_required
def eliminar_reparacion(id):
    db = get_db()
    r = db.execute('SELECT * FROM Reparaciones WHERE id = ?', (id,)).fetchone()
    nombre_eq = r['equipo'] if r else 'Unknown'
    if not r:
        flash('La reparación ya no existe.', 'warning')
        return redirect(url_for('index'))

    # Mantiene una única ventana breve de deshacer por usuario.
    undo_delete_cache[session['user_id']] = {
        'timestamp': time.time(),
        'repair': dict(r),
    }
    db.execute('DELETE FROM Reparaciones WHERE id = ?', (id,))
    db.commit()
    registrar_historial(session['user_id'], session['username'], 'ELIMINAR', 'Reparacion', id, f'Eliminar: {nombre_eq}')
    undo_until = int((time.time() + 5) * 1000)
    return redirect(url_for('index', deleted=1, undo_until=undo_until))

@app.route('/reparacion/restaurar-ultima', methods=['POST'])
@admin_required
def restaurar_ultima_reparacion():
    cached = undo_delete_cache.get(session['user_id'])
    if not cached:
        flash('No hay una eliminación reciente para deshacer.', 'warning')
        return redirect(url_for('index'))

    if time.time() - cached['timestamp'] > 5:
        undo_delete_cache.pop(session['user_id'], None)
        flash('La ventana para deshacer ya expiró.', 'warning')
        return redirect(url_for('index'))

    repair = cached['repair']
    db = get_db()
    db.execute('''
        INSERT INTO Reparaciones (
            id, fecha, sala, uid, npu, parte, familia, equipo, urgente,
            tecnico_id, estado, observaciones, dia_semana,
            fecha_en_reparacion, fecha_reparado, fecha_pendiente, fecha_sin_reparacion
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    ''', (
        repair['id'], repair['fecha'], repair['sala'], repair['uid'], repair['npu'], repair['parte'],
        repair['familia'], repair['equipo'], repair['urgente'], repair['tecnico_id'], repair['estado'],
        repair['observaciones'], repair['dia_semana'], repair['fecha_en_reparacion'],
        repair['fecha_reparado'], repair['fecha_pendiente'], repair['fecha_sin_reparacion']
    ))
    db.commit()
    undo_delete_cache.pop(session['user_id'], None)
    registrar_historial(session['user_id'], session['username'], 'CREAR', 'Reparacion', repair['id'], f'Restaurar: {repair["equipo"]}')
    flash('Reparación restaurada correctamente.', 'success')
    return redirect(url_for('index'))

# =====================================================================
# APIS PARA LA FICHA TÉCNICA
# =====================================================================
@app.route('/api/reparacion/<int:id>/comentario', methods=['POST'])
def agregar_comentario(id):
    db = get_db()
    data = request.get_json()
    nuevo_comentario = data.get('comentario', '').strip()
    if not nuevo_comentario:
        return jsonify({'success': False, 'error': 'El comentario no puede estar vacío.'})
        
    tecnico_id = session.get('tecnico_id')
    
    if tecnico_id:
        tec = db.execute('SELECT nombre FROM Tecnicos WHERE id = ?', (tecnico_id,)).fetchone()
        usuario_actual = tec['nombre'] if tec else session.get('username', 'Usuario')
    else:
        usuario_actual = session.get('username', 'Usuario')

    timestamp = datetime.datetime.now().strftime('%d/%m/%Y %H:%M')
    texto_agregado = f"[{timestamp}] - {usuario_actual}:\n{nuevo_comentario}"
    
    rep = db.execute('SELECT observaciones FROM Reparaciones WHERE id = ?', (id,)).fetchone()
    obs_actual = rep['observaciones'] or ''
    nueva_obs = obs_actual.strip() + "\n\n" + texto_agregado if obs_actual.strip() else texto_agregado
    
    db.execute('UPDATE Reparaciones SET observaciones = ? WHERE id = ?', (nueva_obs, id))
    db.commit()
    registrar_historial(session.get('user_id'), session.get('username', 'unknown'), 'COMENTARIO', 'Reparacion', id, f'Comentario agregado por {usuario_actual}')
    
    return jsonify({'success': True, 'timestamp': timestamp, 'texto': f"{usuario_actual}:\n{nuevo_comentario}"})

@app.route('/api/reparacion/<int:id>/tecnico', methods=['POST'])
def cambiar_tecnico_rapido(id):
    db = get_db()
    data = request.get_json()
    nuevo_tecnico_id = data.get('tecnico_id')
    nuevo_tecnico_id = int(nuevo_tecnico_id) if nuevo_tecnico_id else None

    reparacion = db.execute('SELECT estado FROM Reparaciones WHERE id = ?', (id,)).fetchone()
    estado_actual = (reparacion['estado'] if reparacion else '').upper()
    if session.get('rol') == 'tecnico' and 'REPARADO' in estado_actual:
        return jsonify({'success': False, 'error': 'La reparación ya está cerrada. Solo puedes agregar comentarios.'}), 403

    timestamp = datetime.datetime.now().strftime('%d/%m/%Y %H:%M')
    
    if nuevo_tecnico_id:
        tec = db.execute('SELECT nombre FROM Tecnicos WHERE id = ?', (nuevo_tecnico_id,)).fetchone()
        nombre_tec = tec['nombre'] if tec else 'Desconocido'
        texto_agregado = f"[{timestamp}] - SISTEMA:\nSe asignó el técnico: {nombre_tec}."
    else:
        texto_agregado = f"[{timestamp}] - SISTEMA:\nSe quitó la asignación de técnico (Pasó a Sin Asignar)."

    rep = db.execute('SELECT observaciones FROM Reparaciones WHERE id = ?', (id,)).fetchone()
    obs_actual = rep['observaciones'] or ''
    nueva_obs = obs_actual.strip() + "\n\n" + texto_agregado if obs_actual.strip() else texto_agregado

    db.execute('UPDATE Reparaciones SET tecnico_id = ?, observaciones = ? WHERE id = ?', (nuevo_tecnico_id, nueva_obs, id))
    db.commit()

    nombre_tec = db.execute('SELECT nombre FROM Tecnicos WHERE id = ?', (nuevo_tecnico_id,)).fetchone()['nombre'] if nuevo_tecnico_id else 'Sin asignar'
    registrar_historial(session.get('user_id'), session.get('username', 'unknown'), 'ASIGNAR_TECNICO', 'Reparacion', id, f'Técnico: {nombre_tec}')

    texto_corto = f"SISTEMA:\nSe asignó el técnico: {nombre_tec}." if nuevo_tecnico_id else "SISTEMA:\nSe quitó la asignación de técnico."
    return jsonify({'success': True, 'timestamp': timestamp, 'texto': texto_corto})

@app.route('/api/reparacion/<int:id>/estado', methods=['POST'])
def cambiar_estado_rapido(id):
    db = get_db()
    data = request.get_json()
    nuevo_estado = data.get('estado', '').strip()
    comentario_estado = data.get('comentario', '').strip()

    reparacion = db.execute('SELECT estado FROM Reparaciones WHERE id = ?', (id,)).fetchone()
    estado_actual = (reparacion['estado'] if reparacion else '').upper()
    if session.get('rol') == 'tecnico' and 'REPARADO' in estado_actual:
        return jsonify({'success': False, 'error': 'La reparación ya está cerrada. Solo puedes agregar comentarios.'}), 403
    
    if not nuevo_estado: return jsonify({'success': False})

    estado_upper = nuevo_estado.upper()
    requiere_comentario = ('REPARADO' in estado_upper) or ('PEND' in estado_upper) or ('SIN REPARACION' in estado_upper)
    if requiere_comentario and not comentario_estado:
        return jsonify({'success': False, 'error': 'Debes agregar un comentario para ese cambio de estado.'}), 400

    timestamp = datetime.datetime.now().strftime('%d/%m/%Y %H:%M')
    texto_agregado = f"[{timestamp}] - SISTEMA:\nCambio de estado a: {nuevo_estado}."
    if comentario_estado:
        texto_agregado += f"\nMotivo/Detalle: {comentario_estado}"
    
    rep = db.execute('SELECT observaciones FROM Reparaciones WHERE id = ?', (id,)).fetchone()
    obs_actual = rep['observaciones'] or ''
    nueva_obs = obs_actual.strip() + "\n\n" + texto_agregado if obs_actual.strip() else texto_agregado

    set_clause = "estado = ?, observaciones = ?"
    params = [nuevo_estado, nueva_obs]

    if estado_upper == 'PEND. DE REVISION':
        set_clause += ", fecha_pendiente = ?"
        params.append(timestamp)
    elif estado_upper == 'EN REPARACION':
        set_clause += ", fecha_en_reparacion = ?"
        params.append(timestamp)
    elif 'REPARADO' in estado_upper:
        set_clause += ", fecha_reparado = ?"
        params.append(timestamp)
    elif 'SIN REPARACION' in estado_upper:
        set_clause += ", fecha_sin_reparacion = ?"
        params.append(timestamp)

    params.append(id)
    db.execute(f'UPDATE Reparaciones SET {set_clause} WHERE id = ?', params)
    db.commit()

    rep = db.execute('SELECT equipo FROM Reparaciones WHERE id = ?', (id,)).fetchone()
    nombre_eq = rep['equipo'] if rep else 'Unknown'
    registrar_historial(session.get('user_id'), session.get('username', 'unknown'), 'CAMBIO_ESTADO', 'Reparacion', id, f'{nombre_eq}: {nuevo_estado}')

    texto_respuesta = f"SISTEMA:\nCambio de estado a: {nuevo_estado}."
    if comentario_estado:
        texto_respuesta += f"\nMotivo/Detalle: {comentario_estado}"
    return jsonify({'success': True, 'timestamp': timestamp, 'texto': texto_respuesta})

@app.route('/api/reparacion/<int:id>/prioridad', methods=['POST'])
def cambiar_prioridad_rapido(id):
    db = get_db()
    data = request.get_json()
    urgente = data.get('urgente', 'NO').upper()

    reparacion = db.execute('SELECT estado FROM Reparaciones WHERE id = ?', (id,)).fetchone()
    estado_actual = (reparacion['estado'] if reparacion else '').upper()
    if session.get('rol') == 'tecnico' and 'REPARADO' in estado_actual:
        return jsonify({'success': False, 'error': 'La reparación ya está cerrada. Solo puedes agregar comentarios.'}), 403

    timestamp = datetime.datetime.now().strftime('%d/%m/%Y %H:%M')
    texto_urg = "URGENTE" if urgente == 'SI' else "NORMAL"
    texto_agregado = f"[{timestamp}] - SISTEMA:\nCambio de prioridad a: {texto_urg}."
    
    rep = db.execute('SELECT observaciones FROM Reparaciones WHERE id = ?', (id,)).fetchone()
    obs_actual = rep['observaciones'] or ''
    nueva_obs = obs_actual.strip() + "\n\n" + texto_agregado if obs_actual.strip() else texto_agregado

    db.execute('UPDATE Reparaciones SET urgente = ?, observaciones = ? WHERE id = ?', (urgente, nueva_obs, id))
    db.commit()
    
    return jsonify({'success': True, 'timestamp': timestamp, 'texto': f"SISTEMA:\nCambio de prioridad a: {texto_urg}."})

@app.route('/api/reparacion/<int:id>/devolver', methods=['POST'])
def devolver_reparacion(id):
    db = get_db()
    data = request.get_json()
    motivo = data.get('comentario', '').strip()

    reparacion = db.execute('SELECT estado FROM Reparaciones WHERE id = ?', (id,)).fetchone()
    estado_actual = (reparacion['estado'] if reparacion else '').upper()
    if session.get('rol') == 'tecnico' and 'REPARADO' in estado_actual:
        return jsonify({'success': False, 'error': 'La reparación ya está cerrada. Solo puedes agregar comentarios.'}), 403

    if not motivo:
        return jsonify({'success': False, 'error': 'El motivo de devolución es obligatorio.'})

    timestamp = datetime.datetime.now().strftime('%d/%m/%Y %H:%M')
    
    db.execute('''
        UPDATE Reparaciones 
        SET tecnico_id = NULL, estado = 'PEND. DE REVISION', fecha_pendiente = ? 
        WHERE id = ?
    ''', (timestamp, id))

    texto_agregado = f"[{timestamp}] - 🔙 DEVOLUCIÓN DE EQUIPO\nMotivo: {motivo}\n(El sistema desasignó al técnico y volvió el estado a Pendiente de Revisión)"
    
    rep = db.execute('SELECT observaciones FROM Reparaciones WHERE id = ?', (id,)).fetchone()
    obs_actual = rep['observaciones'] or ''
    nueva_obs = obs_actual.strip() + "\n\n" + texto_agregado if obs_actual.strip() else texto_agregado
    
    db.execute('UPDATE Reparaciones SET observaciones = ? WHERE id = ?', (nueva_obs, id))
    db.commit()

    rep = db.execute('SELECT equipo FROM Reparaciones WHERE id = ?', (id,)).fetchone()
    nombre_eq = rep['equipo'] if rep else 'Unknown'
    registrar_historial(session.get('user_id'), session.get('username', 'unknown'), 'DEVOLVER', 'Reparacion', id, f'{nombre_eq}: Devuelta - {motivo}')

    return jsonify({
        'success': True,
        'timestamp': timestamp,
        'texto': f"🔙 DEVOLUCIÓN DE EQUIPO\nMotivo: {motivo}\n(Estado actualizado a PEND. DE REVISION)"
    })

# =====================================================================
# RUTAS DE EXPORTACIÓN (NUEVO)
# =====================================================================
@app.route('/exportar/excel')
def exportar_excel():
    db = get_db()
    
    tabs_estado = list(INDEX_FILTERS)
    is_tecnico = session.get('rol') == 'tecnico'
    if is_tecnico:
        tabs_estado.insert(1, ('MIS_REPARACIONES', 'Mis rep.'))
        
    estados_validos = {estado for estado, _ in tabs_estado}
    estado_defecto = 'MIS_REPARACIONES' if is_tecnico else 'PEND_REPARACION'
    estado_actual = request.args.get('estado', estado_defecto, type=str).upper()
    if estado_actual not in estados_validos: 
        estado_actual = estado_defecto

    busqueda = request.args.get('q', '', type=str).strip()
    f_sala, f_tecnico, f_estado = request.args.get('f_sala', ''), request.args.get('f_tecnico', ''), request.args.get('f_estado', '')
    f_anio, f_mes, f_dia, f_diasemana = request.args.get('f_anio', ''), request.args.get('f_mes', ''), request.args.get('f_dia', ''), request.args.get('f_diasemana', '')

    base_where_clauses, base_query_params = [], []
    if f_sala: base_where_clauses.append('r.sala = ?'); base_query_params.append(f_sala)
    if f_tecnico:
        if f_tecnico == 'SIN_ASIGNAR': base_where_clauses.append('r.tecnico_id IS NULL')
        else: base_where_clauses.append('t.nombre = ?'); base_query_params.append(f_tecnico)
    if f_estado: base_where_clauses.append('r.estado = ?'); base_query_params.append(f_estado)
    if f_anio: base_where_clauses.append('r.fecha LIKE ?'); base_query_params.append(f'%/{f_anio}%')
    if f_mes: base_where_clauses.append('(r.fecha LIKE ? OR r.fecha LIKE ?)'); base_query_params.extend([f'%/{int(f_mes):02d}/%', f'%/{int(f_mes)}/%'])
    if f_dia: base_where_clauses.append('(r.fecha LIKE ? OR r.fecha LIKE ?)'); base_query_params.extend([f'{int(f_dia):02d}/%', f'{int(f_dia)}/%'])
    if f_diasemana: base_where_clauses.append('r.dia_semana = ?'); base_query_params.append(int(f_diasemana))
    if busqueda:
        base_where_clauses.append('(' + ' OR '.join(['r.fecha LIKE ?', 'r.sala LIKE ?', 'r.uid LIKE ?', 'r.npu LIKE ?', 'r.equipo LIKE ?', 'r.estado LIKE ?', 'r.observaciones LIKE ?', 't.nombre LIKE ?']) + ')')
        base_query_params.extend([f'%{busqueda}%'] * 8)

    final_where_clauses, final_query_params = list(base_where_clauses), list(base_query_params)
    tecnico_id_actual = session.get('tecnico_id') or 0
    
    if estado_actual == 'URGENTES': 
        final_where_clauses.extend(["r.urgente = 'SI'", "UPPER(r.estado) NOT LIKE 'REPARADO%'", "UPPER(r.estado) NOT LIKE 'SIN REPARACION%'"])
    elif estado_actual == 'MIS_REPARACIONES': 
        final_where_clauses.extend(["r.tecnico_id = ?", "UPPER(r.estado) NOT LIKE 'REPARADO%'", "UPPER(r.estado) NOT LIKE 'SIN REPARACION%'"])
        final_query_params.append(tecnico_id_actual)
    elif estado_actual == 'PEND_REPARACION': 
        final_where_clauses.append('UPPER(r.estado) = ?'); final_query_params.append('PEND. DE REVISION')
    elif estado_actual == 'EN_REPARACION': 
        final_where_clauses.append('UPPER(r.estado) = ?'); final_query_params.append('EN REPARACION')
    elif estado_actual == 'REPARADOS': 
        final_where_clauses.append('UPPER(r.estado) LIKE ?'); final_query_params.append('REPARADO%')
    elif estado_actual == 'SIN_REPARACION': 
        final_where_clauses.append('UPPER(r.estado) LIKE ?'); final_query_params.append('SIN REPARACION%')
    elif estado_actual == 'PENDIENTES': 
        final_where_clauses.extend(["UPPER(r.estado) NOT LIKE 'REPARADO%'", "UPPER(r.estado) NOT LIKE 'SIN REPARACION%'", "UPPER(r.estado) != 'PEND. DE REVISION'", "UPPER(r.estado) != 'EN REPARACION'"])

    final_where_sql = (' WHERE ' + ' AND '.join(final_where_clauses)) if final_where_clauses else ''
    
    query = f'''
        SELECT r.id, r.fecha, r.sala, r.uid, r.npu, r.familia, r.equipo, r.urgente, 
               r.estado, r.observaciones, r.fecha_en_reparacion, r.fecha_reparado, 
               t.nombre AS tecnico_nombre 
        FROM Reparaciones r 
        LEFT JOIN Tecnicos t ON r.tecnico_id = t.id 
        {final_where_sql} 
        ORDER BY r.id DESC
    '''
    reparaciones = db.execute(query, final_query_params).fetchall()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reparaciones"

    headers = ['ID', 'Fecha Ingreso', 'Sala', 'UID', 'NPU', 'Familia', 'Equipo', 'Urgente', 'Técnico', 'Estado', 'Fecha En Rep.', 'Fecha Reparado', 'Observaciones']
    ws.append(headers)
    
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for r in reparaciones:
        ws.append([
            r['id'], r['fecha'], r['sala'], r['uid'], r['npu'], r['familia'], 
            r['equipo'], r['urgente'], r['tecnico_nombre'] or 'Sin asignar', 
            r['estado'], r['fecha_en_reparacion'], r['fecha_reparado'], r['observaciones']
        ])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    timestamp = datetime.datetime.now().strftime('%Y%m%d_%H%M')
    filename = f'Reporte_{estado_actual}_{timestamp}.xlsx'

    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

# =====================================================================
# DASHBOARD DE ESTADÍSTICAS (KPIs)
# =====================================================================
@app.route('/dashboard')
@admin_required
def dashboard():
    db = get_db()
    
    # Filtros de entrada
    f_anio = request.args.get('f_anio', datetime.datetime.now().year, type=int)
    f_sala = request.args.get('f_sala', '')
    
    # Opciones para los selectores de filtro
    anios_opt = db.execute("SELECT DISTINCT SUBSTR(fecha, 7, 4) as anio FROM Reparaciones WHERE fecha LIKE '__/__/____%' ORDER BY anio DESC").fetchall()
    if not anios_opt: anios_opt = [{'anio': datetime.datetime.now().year}]
    
    salas_opt = db.execute("SELECT nombre as sala FROM SalasCatalogo ORDER BY nombre").fetchall()
    
    # --- 1. EQUIPOS QUE MÁS FALLAN (Basado en fecha de ingreso) ---
    where_fails = ["r.fecha LIKE ?"]
    params_fails = [f'%/%/{f_anio}%']
    if f_sala:
        where_fails.append("r.sala = ?")
        params_fails.append(f_sala)
    
    where_fails_sql = " WHERE " + " AND ".join(where_fails)
    
    # Top Familias
    top_familias = db.execute(f'''
        SELECT familia, COUNT(*) as total 
        FROM Reparaciones r
        {where_fails_sql} AND familia IS NOT NULL AND familia != ''
        GROUP BY familia ORDER BY total DESC LIMIT 8
    ''', params_fails).fetchall()
    
    # Top Equipos
    top_equipos = db.execute(f'''
        SELECT equipo, COUNT(*) as total 
        FROM Reparaciones r
        {where_fails_sql}
        GROUP BY equipo ORDER BY total DESC LIMIT 10
    ''', params_fails).fetchall()
    
    # --- 2. PRODUCTIVIDAD (Basado en fecha de reparado) ---
    where_prod = ["r.fecha_reparado LIKE ?"]
    params_prod = [f'%/%/{f_anio}%']
    if f_sala:
        where_prod.append("r.sala = ?")
        params_prod.append(f_sala)
    
    where_prod_sql = " WHERE " + " AND ".join(where_prod) + " AND UPPER(r.estado) LIKE 'REPARADO%'"
    
    # Por Técnico
    prod_tecnico = db.execute(f'''
        SELECT t.nombre, COUNT(*) as total 
        FROM Reparaciones r
        JOIN Tecnicos t ON r.tecnico_id = t.id
        {where_prod_sql}
        GROUP BY t.nombre ORDER BY total DESC
    ''', params_prod).fetchall()
    
    # Por Sala
    prod_sala = db.execute(f'''
        SELECT r.sala, COUNT(*) as total 
        FROM Reparaciones r
        {where_prod_sql}
        GROUP BY r.sala ORDER BY total DESC LIMIT 10
    ''', params_prod).fetchall()
    
    # --- 3. EVOLUCIÓN TEMPORAL (Mensual) ---
    prod_mensual_raw = db.execute(f'''
        SELECT SUBSTR(r.fecha_reparado, 4, 2) as mes, COUNT(*) as total
        FROM Reparaciones r
        {where_prod_sql}
        GROUP BY mes ORDER BY mes ASC
    ''', params_prod).fetchall()
    
    meses_nombres = ['Ene', 'Feb', 'Mar', 'Abr', 'May', 'Jun', 'Jul', 'Ago', 'Sep', 'Oct', 'Nov', 'Dic']
    data_mes = [0] * 12
    for row in prod_mensual_raw:
        try:
            idx = int(row['mes']) - 1
            if 0 <= idx < 12: data_mes[idx] = row['total']
        except: continue

    # --- 4. KPIs RÁPIDOS (Hoy, Semana, Mes, Año - Solo Reparados) ---
    where_kpi = ["UPPER(estado) LIKE 'REPARADO%'"]
    params_kpi = []
    if f_sala:
        where_kpi.append("sala = ?")
        params_kpi.append(f_sala)
    
    where_kpi_sql = " AND ".join(where_kpi)
    
    hoy_str = datetime.datetime.now().strftime('%d/%m/%Y')
    kpi_hoy = db.execute(f"SELECT COUNT(*) as total FROM Reparaciones WHERE fecha_reparado LIKE ? AND {where_kpi_sql}", [f'{hoy_str}%'] + params_kpi).fetchone()['total']
    
    siete_dias_atras = (datetime.datetime.now() - datetime.timedelta(days=7)).strftime('%Y-%m-%d')
    kpi_semana = db.execute(f'''
        SELECT COUNT(*) as total FROM Reparaciones 
        WHERE (SUBSTR(fecha_reparado, 7, 4) || '-' || SUBSTR(fecha_reparado, 4, 2) || '-' || SUBSTR(fecha_reparado, 1, 2)) >= ?
        AND {where_kpi_sql}
    ''', [siete_dias_atras] + params_kpi).fetchone()['total']
    
    este_mes_str = datetime.datetime.now().strftime('/%m/%Y')
    kpi_mes = db.execute(f"SELECT COUNT(*) as total FROM Reparaciones WHERE fecha_reparado LIKE ? AND {where_kpi_sql}", [f'%{este_mes_str}%'] + params_kpi).fetchone()['total']
    
    kpi_anio = db.execute(f"SELECT COUNT(*) as total FROM Reparaciones WHERE fecha_reparado LIKE ? AND {where_kpi_sql}", [f'%/%/{f_anio}%'] + params_kpi).fetchone()['total']

    return render_template('dashboard.html',
        f_anio=f_anio, f_sala=f_sala,
        anios_opt=anios_opt, salas_opt=salas_opt,
        top_familias=top_familias, top_equipos=top_equipos,
        prod_tecnico=prod_tecnico, prod_sala=prod_sala,
        labels_mes=meses_nombres, data_mes=data_mes,
        kpis={'hoy': kpi_hoy, 'semana': kpi_semana, 'mes': kpi_mes, 'anio': kpi_anio}
    )

# =====================================================================
# CONFIGURACIÓN Y ADMINISTRACIÓN
# =====================================================================
@app.route('/configuracion')
@admin_required
def configuracion():
    db = get_db()
    catalogos = {tipo: get_catalog_items(db, tipo) for tipo in CATALOGS}
    usuarios = db.execute('''
        SELECT u.id, u.username, u.rol, t.nombre as tecnico_nombre 
        FROM Usuarios u 
        LEFT JOIN Tecnicos t ON u.tecnico_id = t.id 
        ORDER BY u.rol, u.username
    ''').fetchall()
    reparaciones_raw = db.execute('SELECT * FROM Reparaciones ORDER BY id DESC').fetchall()
    historial_reciente = db.execute('SELECT * FROM Historial ORDER BY fecha DESC LIMIT 15').fetchall()
    os.makedirs('backups', exist_ok=True)
    backups = []
    for name in sorted(os.listdir('backups'), reverse=True):
        if not name.lower().endswith('.db'):
            continue
        path = os.path.join('backups', name)
        stat = os.stat(path)
        backups.append({
            'filename': name,
            'size_kb': max(1, int(stat.st_size / 1024)),
            'modified': datetime.datetime.fromtimestamp(stat.st_mtime).strftime('%d/%m/%Y %H:%M')
        })
    return render_template(
        'configuracion.html',
        catalogos=catalogos,
        catalog_config=CATALOGS,
        usuarios=usuarios,
        reparaciones_raw=reparaciones_raw,
        historial_reciente=historial_reciente,
        backups=backups[:20]
    )

@app.route('/configuracion/usuarios/agregar', methods=['POST'])
@admin_required
def agregar_usuario():
    username = request.form.get('username', '').strip()
    password = request.form.get('password', '')
    rol = request.form.get('rol')
    tecnico_id = request.form.get('tecnico_id') if rol == 'tecnico' else None
    
    if not username or not password:
        flash('El usuario y la contraseña son obligatorios.', 'warning')
        return redirect(url_for('configuracion') + '#panel-usuarios')

    db = get_db()
    try:
        db.execute('INSERT INTO Usuarios (username, password_hash, rol, tecnico_id) VALUES (?, ?, ?, ?)',
                   (username, generate_password_hash(password), rol, tecnico_id))
        db.commit()
        registrar_historial(session['user_id'], session['username'], 'CREAR', 'Usuario', None, f'Crear usuario: {username} ({rol})')
        flash('Usuario creado exitosamente.', 'success')
    except sqlite3.IntegrityError:
        flash('El nombre de usuario ya existe. Elegí otro.', 'danger')

    return redirect(url_for('configuracion') + '#panel-usuarios')

@app.route('/configuracion/usuarios/editar/<int:id>', methods=['POST'])
@admin_required
def editar_usuario(id):
    db = get_db()
    username = request.form.get('username', '').strip()
    password = request.form.get('password', '')
    rol = request.form.get('rol')
    tecnico_id = request.form.get('tecnico_id') if rol == 'tecnico' else None

    if not username:
        flash('El nombre de usuario es obligatorio.', 'warning')
        return redirect(url_for('configuracion') + '#panel-usuarios')

    try:
        if password:
            nueva_pass_hash = generate_password_hash(password)
            db.execute('''
                UPDATE Usuarios 
                SET username = ?, password_hash = ?, rol = ?, tecnico_id = ? 
                WHERE id = ?
            ''', (username, nueva_pass_hash, rol, tecnico_id, id))
        else:
            db.execute('''
                UPDATE Usuarios
                SET username = ?, rol = ?, tecnico_id = ?
                WHERE id = ?
            ''', (username, rol, tecnico_id, id))

        db.commit()
        registrar_historial(session['user_id'], session['username'], 'EDITAR', 'Usuario', id, f'Editar usuario: {username}')
        flash('Usuario actualizado correctamente.', 'success')
    except sqlite3.IntegrityError:
        flash('El nombre de usuario ya está en uso.', 'danger')

    return redirect(url_for('configuracion') + '#panel-usuarios')

@app.route('/configuracion/usuarios/eliminar/<int:id>', methods=['POST'])
@admin_required
def eliminar_usuario(id):
    db = get_db()
    user = db.execute('SELECT username FROM Usuarios WHERE id = ?', (id,)).fetchone()
    if user and user['username'] == 'admin':
        flash('No se puede eliminar al administrador principal.', 'danger')
    else:
        username = user['username'] if user else 'Unknown'
        db.execute('DELETE FROM Usuarios WHERE id = ?', (id,))
        db.commit()
        registrar_historial(session['user_id'], session['username'], 'ELIMINAR', 'Usuario', id, f'Eliminar usuario: {username}')
        flash('Usuario eliminado.', 'success')
    return redirect(url_for('configuracion') + '#panel-usuarios')

@app.route('/configuracion/agregar/<tipo>', methods=['POST'])
@admin_required
def agregar_catalogo(tipo):
    if tipo not in CATALOGS: return redirect(url_for('configuracion'))
    db = get_db()
    nombre = request.form.get('nombre', '').strip().upper()
    
    if tipo == 'equipos':
        valor = normalize_price_input(request.form.get('valor', ''))
        lab = request.form.get('lab', '').strip().upper()
        nombre_final = f"{nombre} - {lab}" if lab else nombre
        db.execute('INSERT INTO EquiposCatalogo (nombre, valor) VALUES (?, ?)', (nombre_final, valor))
    elif tipo == 'tecnicos':
        db.execute('INSERT INTO Tecnicos (nombre, activo) VALUES (?, 1)', (nombre,))
    else:
        table = CATALOGS[tipo]['table']
        db.execute(f'INSERT INTO {table} (nombre) VALUES (?)', (nombre,))

    db.commit()
    registrar_historial(session['user_id'], session['username'], 'CREAR', tipo.capitalize(), None, f'Crear {tipo}: {nombre}')
    flash(f'Elemento agregado a {CATALOGS[tipo]["title"]}.', 'success')
    return redirect(url_for('configuracion') + f'#panel-{tipo}')

@app.route('/configuracion/editar/<tipo>/<int:item_id>', methods=['POST'])
@admin_required
def editar_catalogo(tipo, item_id):
    if tipo not in CATALOGS: return redirect(url_for('configuracion'))
    db = get_db()
    nombre = request.form.get('nombre', '').strip().upper()
    
    if tipo == 'equipos':
        valor = normalize_price_input(request.form.get('valor', ''))
        db.execute('UPDATE EquiposCatalogo SET nombre = ?, valor = ? WHERE id = ?', (nombre, valor, item_id))
    elif tipo == 'tecnicos':
        db.execute('UPDATE Tecnicos SET nombre = ? WHERE id = ?', (nombre, item_id))
    else:
        table = CATALOGS[tipo]['table']
        db.execute(f'UPDATE {table} SET nombre = ? WHERE id = ?', (nombre, item_id))

    db.commit()
    registrar_historial(session['user_id'], session['username'], 'EDITAR', tipo.capitalize(), item_id, f'Editar {tipo}: {nombre}')
    flash('Elemento editado correctamente.', 'success')
    return redirect(url_for('configuracion') + f'#panel-{tipo}')

@app.route('/configuracion/eliminar/<tipo>/<int:item_id>', methods=['GET', 'POST'])
@admin_required
def eliminar_catalogo(tipo, item_id):
    if tipo not in CATALOGS: return redirect(url_for('configuracion'))
    db = get_db()
    table = CATALOGS[tipo]['table']
    try:
        item = db.execute(f'SELECT nombre FROM {table} WHERE id = ?', (item_id,)).fetchone()
        nombre_item = item['nombre'] if item else 'Unknown'
        db.execute(f'DELETE FROM {table} WHERE id = ?', (item_id,))
        db.commit()
        registrar_historial(session['user_id'], session['username'], 'ELIMINAR', tipo.capitalize(), item_id, f'Eliminar {tipo}: {nombre_item}')
        flash('Elemento eliminado.', 'success')
    except sqlite3.IntegrityError:
        db.rollback()
        flash(f'No se puede eliminar: Este elemento está siendo usado en una o más reparaciones.', 'danger')
    return redirect(url_for('configuracion') + f'#panel-{tipo}')

@app.route('/configuracion/tecnico/desactivar/<int:item_id>', methods=['GET', 'POST'])
@admin_required
def desactivar_tecnico(item_id):
    db = get_db()
    db.execute('UPDATE Tecnicos SET activo = 0 WHERE id = ?', (item_id,))
    db.commit()
    flash('Técnico desactivado.', 'success')
    return redirect(url_for('configuracion') + '#panel-tecnicos')

@app.route('/configuracion/tecnico/reactivar/<int:item_id>', methods=['POST'])
@admin_required
def reactivar_tecnico(item_id):
    db = get_db()
    db.execute('UPDATE Tecnicos SET activo = 1 WHERE id = ?', (item_id,))
    db.commit()
    flash('Técnico reactivado.', 'success')
    return redirect(url_for('configuracion') + '#panel-tecnicos')

@app.route('/configuracion/tecnico/eliminar_permanente/<int:item_id>', methods=['GET', 'POST'])
@admin_required
def eliminar_tecnico_permanente(item_id):
    db = get_db()
    try:
        db.execute('DELETE FROM Tecnicos WHERE id = ?', (item_id,))
        db.commit()
        flash('Técnico eliminado permanentemente.', 'success')
        return redirect(url_for('configuracion') + '#panel-tecnicos')
    except sqlite3.IntegrityError:
        db.rollback()
        return redirect(url_for('configuracion', error_tecnico_en_uso=item_id) + '#panel-tecnicos')

@app.route('/configuracion/equipos/ajuste_global', methods=['POST'])
@admin_required
def ajuste_global_precios():
    db = get_db()
    operacion = request.form.get('operacion')
    tipo_ajuste = request.form.get('tipo_ajuste')
    try:
        valor_ajuste = float(request.form.get('valor_ajuste', 0))
    except:
        flash('Valor de ajuste inválido.', 'danger')
        return redirect(url_for('configuracion') + '#panel-equipos')

    equipos = db.execute('SELECT id, valor FROM EquiposCatalogo').fetchall()
    for eq in equipos:
        try:
            precio_actual = stored_price_to_amount(eq['valor']) if eq['valor'] else 0.0
        except:
            precio_actual = 0.0
            
        if tipo_ajuste == 'porcentaje':
            modificacion = precio_actual * (valor_ajuste / 100.0)
        else:
            modificacion = valor_ajuste

        nuevo_precio = precio_actual + modificacion if operacion == 'aumento' else precio_actual - modificacion
        if nuevo_precio < 0: nuevo_precio = 0
        db.execute('UPDATE EquiposCatalogo SET valor = ? WHERE id = ?', (normalize_price_input(f"{nuevo_precio:.2f}"), eq['id']))

    db.commit()
    flash('Ajuste de precios aplicado a todos los equipos.', 'success')
    return redirect(url_for('configuracion') + '#panel-equipos')

@app.route('/configuracion/equipos/acciones_masivas', methods=['POST'])
@admin_required
def acciones_masivas_equipos():
    accion = request.form.get('accion')
    ids = request.form.getlist('ids[]')
    if not ids:
        flash('No seleccionaste ningún equipo.', 'warning')
        return redirect(url_for('configuracion') + '#panel-equipos')
        
    db = get_db()
    if accion == 'eliminar':
        placeholders = ','.join('?' * len(ids))
        db.execute(f'DELETE FROM EquiposCatalogo WHERE id IN ({placeholders})', ids)
        db.commit()
        flash(f'Se eliminaron {len(ids)} equipos.', 'success')
        
    return redirect(url_for('configuracion') + '#panel-equipos')

@app.route('/configuracion/salas/sincronizar', methods=['POST'])
@admin_required
def sincronizar_salas():
    db = get_db()
    salas_unicas = db.execute('SELECT DISTINCT sala FROM Reparaciones WHERE sala IS NOT NULL AND sala != ""').fetchall()
    for s in salas_unicas:
        ensure_catalog_value(db, 'SalasCatalogo', s['sala'])
    db.commit()
    flash('Salas sincronizadas con la base de reparaciones.', 'success')
    return redirect(url_for('configuracion') + '#panel-salas')

@app.route('/configuracion/recalcular_dias', methods=['POST'])
@admin_required
def recalcular_dias():
    db = get_db()
    reps = db.execute('SELECT id, fecha FROM Reparaciones').fetchall()
    for r in reps:
        dia = calcular_dia_semana(r['fecha'])
        if dia is not None:
            db.execute('UPDATE Reparaciones SET dia_semana = ? WHERE id = ?', (dia, r['id']))
    db.commit()
    flash('Fechas analizadas y días de la semana recalculados.', 'success')
    return redirect(url_for('configuracion') + '#panel-reparaciones')

@app.route('/configuracion/vaciar_reparaciones', methods=['POST'])
@admin_required
def vaciar_reparaciones():
    db = get_db()
    db.execute('DELETE FROM Reparaciones')
    db.commit()
    flash('Toda la base de reparaciones ha sido vaciada.', 'danger')
    return redirect(url_for('configuracion') + '#panel-reparaciones')

@app.route('/api/reparaciones/preview_csv', methods=['POST'])
@admin_required
def preview_reparaciones_csv():
    try:
        archivo = request.files.get('archivo_csv')
        if not archivo or not archivo.filename:
            return jsonify({'error': 'Debes seleccionar un archivo CSV.'}), 400
        if not archivo.filename.lower().endswith('.csv'):
            return jsonify({'error': 'El archivo debe tener extension .csv.'}), 400

        try:
            contenido = decode_csv_bytes(archivo.read())
        except ValueError as e:
            return jsonify({'error': str(e)}), 400

        # Detección automática de delimitador
        delimiter = ','
        if contenido:
            first_line = contenido.split('\n')[0]
            if '\t' in first_line: delimiter = '\t'
            elif ';' in first_line: delimiter = ';'
            elif '|' in first_line: delimiter = '|'

        reader = csv.DictReader(io.StringIO(contenido), delimiter=delimiter)
        if not reader.fieldnames:
            return jsonify({'error': 'El CSV no tiene cabecera.'}), 400

        normalized_headers = [h.strip().lower() for h in reader.fieldnames if h and h.strip()]
        
        # Mapeo de columnas requeridas a sus posibles nombres (normalizados)
        column_maps = {
            'fecha': ['fecha'],
            'sala': ['sala'],
            'uid': ['uid'],
            'npu': ['npu'],
            'familia': ['familia'],
            'equipo': ['equipo'],
            'urgente': ['urgente'],
            'tecnico': ['tecnico'],
            'estados de reparacion': ['estados de reparacion', 'estado', 'estados'],
            'observaciones': ['observaciones', 'obs.', 'obs']
        }
        
        missing_headers = []
        for req, aliases in column_maps.items():
            if not any(alias in normalized_headers for alias in aliases):
                missing_headers.append(req)

        if missing_headers:
            return jsonify({'error': 'Cabecera invalida. Faltan columnas: ' + ', '.join(missing_headers), 'missing_headers': missing_headers, 'delimiter_detected': delimiter}), 400

        optional_headers = ['f. pendiente', 'f. en rep.', 'f. reparado', 'f. sin rep.', 'obs.', 
                            'fecha "pendiente"', 'fecha "en reparacion / en prueba "', 'fecha "reparado"', 'fecha "sin reparacion"']
        has_date_columns = any(h in normalized_headers for h in optional_headers)

        reader.fieldnames = normalized_headers
        rows = list(reader)

        total_rows = len(rows)
        valid_rows = 0
        skipped_rows = 0
        preview = []

        for row in rows[:50]:
            fecha = normalize_date(get_csv_value(row, 'fecha'))
            sala = get_csv_value(row, 'sala') or CATALOGS['salas']['default']
            uid = get_csv_value(row, 'uid')
            npu = get_csv_value(row, 'npu')
            familia = get_csv_value(row, 'familia')
            equipo = get_csv_value(row, 'equipo') or CATALOGS['equipos']['default']
            urgente = get_csv_value(row, 'urgente', 'NO').upper()
            tecnico_nombre = get_csv_value(row, 'tecnico', 'SIN ASIGNAR')
            estado = get_csv_value(row, 'estados de reparacion', CATALOGS['estados']['default']).upper() or CATALOGS['estados']['default']
            observaciones = get_csv_value(row, 'observaciones') or get_csv_value(row, 'obs.')

            if has_date_columns:
                fp = normalize_date(get_csv_value(row, 'f. pendiente'))
                fr = normalize_date(get_csv_value(row, 'f. en rep.'))
                frep = normalize_date(get_csv_value(row, 'f. reparado'))
                fsr = normalize_date(get_csv_value(row, 'f. sin rep.'))
                
                # FALLBACK logic
                if not frep and 'REPARADO' in estado: frep = fecha
                if not fsr and 'SIN REPARACION' in estado: fsr = fecha
                if not fr and estado == 'EN REPARACION': fr = fecha
                if not fp and not frep and not fsr and not fr: fp = fecha
            else:
                fp = fecha if estado not in ('EN REPARACION',) and 'REPARADO' not in estado and 'SIN REPARACION' not in estado else ''
                fr = fecha if estado == 'EN REPARACION' else ''
                frep = fecha if 'REPARADO' in estado else ''
                fsr = fecha if 'SIN REPARACION' in estado else ''

            valid = bool(fecha and sala and equipo)
            warnings = []
            if not fecha:
                warnings.append('Sin fecha')
            if not sala:
                warnings.append('Sin sala')
            if not equipo:
                warnings.append('Sin equipo')
            if urgente not in ('SI', 'NO') and urgente:
                warnings.append(f'Urgente invalido: {urgente}')
                urgente = 'NO'

            preview.append({
                'fecha': fecha, 'sala': sala, 'uid': uid, 'npu': npu,
                'familia': familia, 'equipo': equipo, 'urgente': urgente,
                'tecnico': tecnico_nombre, 'estado': estado, 'observaciones': observaciones,
                'fecha_pendiente': fp, 'fecha_en_reparacion': fr,
                'fecha_reparado': frep, 'fecha_sin_reparacion': fsr,
                'valid': valid, 'warnings': warnings
            })

        for row in rows:
            f_val = get_csv_value(row, 'fecha')
            s_val = get_csv_value(row, 'sala') or CATALOGS['salas']['default']
            e_val = get_csv_value(row, 'equipo') or CATALOGS['equipos']['default']
            if f_val and s_val and e_val:
                valid_rows += 1
            else:
                skipped_rows += 1

        display_headers = [
            'fecha', 'sala', 'uid', 'npu', 'familia', 'equipo', 'urgente', 'tecnico', 
            'estados de reparacion', 'f. pendiente', 'f. en rep.', 'f. reparado', 
            'f. sin rep.', 'observaciones'
        ]

        return jsonify({
            'total_rows': total_rows,
            'valid_rows': valid_rows,
            'skipped_rows': skipped_rows,
            'headers': display_headers,
            'has_date_columns': has_date_columns,
            'preview': preview,
            'preview_limit': 50,
            'has_more': total_rows > 50
        })
    except Exception as e:
        print(f"Error en preview_reparaciones_csv: {e}")
        return jsonify({'error': f'Error interno del servidor: {str(e)}'}), 500



@app.route('/configuracion/csv_base', methods=['POST'])
@admin_required
def cargar_csv():
    archivo = request.files.get('archivo_csv')
    if not archivo or not archivo.filename:
        flash('Debes seleccionar un archivo CSV.', 'warning')
        return redirect(url_for('configuracion') + '#panel-reparaciones')

    if not archivo.filename.lower().endswith('.csv'):
        flash('El archivo debe tener extensión .csv.', 'danger')
        return redirect(url_for('configuracion') + '#panel-reparaciones')

    try:
        contenido = decode_csv_bytes(archivo.read())
    except ValueError as e:
        flash(str(e), 'danger')
        return redirect(url_for('configuracion') + '#panel-reparaciones')

    # Detección automática de delimitador
    delimiter = ','
    if contenido:
        first_line = contenido.split('\n')[0]
        if '\t' in first_line: delimiter = '\t'
        elif ';' in first_line: delimiter = ';'
        elif '|' in first_line: delimiter = '|'

    reader = csv.DictReader(io.StringIO(contenido), delimiter=delimiter)
    if not reader.fieldnames:
        flash('El CSV no tiene cabecera.', 'danger')
        return redirect(url_for('configuracion') + '#panel-reparaciones')

    normalized_headers = [h.strip().lower() for h in reader.fieldnames if h and h.strip()]
    
    # Mapeo de columnas requeridas a sus posibles nombres (normalizados)
    column_maps = {
        'fecha': ['fecha'],
        'sala': ['sala'],
        'uid': ['uid'],
        'npu': ['npu'],
        'familia': ['familia'],
        'equipo': ['equipo'],
        'urgente': ['urgente'],
        'tecnico': ['tecnico'],
        'estados de reparacion': ['estados de reparacion', 'estado', 'estados'],
        'observaciones': ['observaciones', 'obs.', 'obs']
    }
    
    missing_headers = []
    for req, aliases in column_maps.items():
        if not any(alias in normalized_headers for alias in aliases):
            missing_headers.append(req)

    if missing_headers:
        flash('Cabecera inválida. Faltan columnas: ' + ', '.join(missing_headers), 'danger')
        return redirect(url_for('configuracion') + '#panel-reparaciones')

    optional_headers = ['f. pendiente', 'f. en rep.', 'f. reparado', 'f. sin rep.', 'obs.', 
                        'fecha "pendiente"', 'fecha "en reparacion / en prueba "', 'fecha "reparado"', 'fecha "sin reparacion"']
    has_date_columns = any(h in normalized_headers for h in optional_headers)

    reader.fieldnames = normalized_headers
    rows = list(reader)
    if not rows:
        flash('El CSV no contiene filas para importar.', 'warning')
        return redirect(url_for('configuracion') + '#panel-reparaciones')

    db = get_db()
    backup_path = create_backup_file()

    try:
        db.execute('DELETE FROM Reparaciones')
        db.execute("DELETE FROM sqlite_sequence WHERE name = 'Reparaciones'")

        imported_count = 0
        for row in rows:
            fecha = normalize_date(get_csv_value(row, 'fecha'))
            sala = get_csv_value(row, 'sala') or CATALOGS['salas']['default']
            uid = get_csv_value(row, 'uid')
            npu = get_csv_value(row, 'npu')
            familia = get_csv_value(row, 'familia')
            equipo = get_csv_value(row, 'equipo') or CATALOGS['equipos']['default']
            urgente = get_csv_value(row, 'urgente', 'NO').upper()
            tecnico_nombre = get_csv_value(row, 'tecnico', 'SIN ASIGNAR')
            estado = get_csv_value(row, 'estados de reparacion', CATALOGS['estados']['default']).upper() or CATALOGS['estados']['default']
            observaciones = get_csv_value(row, 'observaciones') or get_csv_value(row, 'obs.')

            if not fecha or not sala or not equipo:
                continue

            if urgente not in ('SI', 'NO'):
                urgente = 'NO'

            tecnico_id = None
            tecnico_upper = tecnico_nombre.upper()
            if tecnico_nombre and tecnico_upper not in ('SIN ASIGNAR', 'SIN ASIGNACION', '--', 'N/A'):
                tecnico = db.execute('SELECT id FROM Tecnicos WHERE UPPER(nombre) = ?', (tecnico_upper,)).fetchone()
                if tecnico:
                    tecnico_id = tecnico['id']
                    db.execute('UPDATE Tecnicos SET activo = 1 WHERE id = ?', (tecnico_id,))
                else:
                    cur = db.execute('INSERT INTO Tecnicos (nombre, activo) VALUES (?, 1)', (tecnico_nombre,))
                    tecnico_id = cur.lastrowid

            dia_semana = calcular_dia_semana(fecha)

            if has_date_columns:
                fecha_pendiente = normalize_date(get_csv_value(row, 'f. pendiente'))
                fecha_en_reparacion = normalize_date(get_csv_value(row, 'f. en rep.'))
                fecha_reparado = normalize_date(get_csv_value(row, 'f. reparado'))
                fecha_sin_reparacion = normalize_date(get_csv_value(row, 'f. sin rep.'))
                
                # FALLBACK: Si no hay fecha específica para el estado actual, usar la fecha de ingreso
                if not fecha_reparado and 'REPARADO' in estado: fecha_reparado = fecha
                if not fecha_sin_reparacion and 'SIN REPARACION' in estado: fecha_sin_reparacion = fecha
                if not fecha_en_reparacion and estado == 'EN REPARACION': fecha_en_reparacion = fecha
                if not fecha_pendiente and not fecha_reparado and not fecha_sin_reparacion and not fecha_en_reparacion:
                    fecha_pendiente = fecha
            else:
                fecha_pendiente = ''
                fecha_en_reparacion = ''
                fecha_reparado = ''
                fecha_sin_reparacion = ''
                if 'REPARADO' in estado:
                    fecha_reparado = fecha
                elif 'SIN REPARACION' in estado:
                    fecha_sin_reparacion = fecha
                elif estado == 'EN REPARACION':
                    fecha_en_reparacion = fecha
                else:
                    fecha_pendiente = fecha

            db.execute('''
                INSERT INTO Reparaciones (
                    fecha, sala, uid, npu, parte, familia, equipo, urgente, tecnico_id,
                    estado, observaciones, dia_semana, fecha_en_reparacion, fecha_reparado,
                    fecha_pendiente, fecha_sin_reparacion
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ''', (
                fecha, sala, uid, npu, '', familia, equipo, urgente, tecnico_id,
                estado, observaciones, dia_semana, fecha_en_reparacion, fecha_reparado,
                fecha_pendiente, fecha_sin_reparacion
            ))

            ensure_catalog_value(db, 'SalasCatalogo', sala)
            ensure_catalog_value(db, 'FamiliasCatalogo', familia)
            ensure_catalog_value(db, 'EquiposCatalogo', equipo)
            ensure_catalog_value(db, 'EstadosCatalogo', estado)
            imported_count += 1

        db.commit()
        registrar_historial(session['user_id'], session['username'], 'CONFIGURACION', 'Sistema', None, f'Importación CSV base: {imported_count} registros. Backup previo: {backup_path}')
        flash(f'Importación completada: {imported_count} reparaciones cargadas. Backup previo: {backup_path}', 'success')
    except Exception as e:
        db.rollback()
        flash(f'Error durante la importación: {str(e)}', 'danger')

    return redirect(url_for('configuracion') + '#panel-reparaciones')

@app.route('/configuracion/csv_equipos', methods=['POST'])
@admin_required
def cargar_equipos_csv():
    flash('Carga de equipos por CSV en desarrollo.', 'info')
    return redirect(url_for('configuracion') + '#panel-equipos')

@app.route('/configuracion/csv_familias', methods=['POST'])
@admin_required
def cargar_familias_csv():
    flash('Carga de familias por CSV en desarrollo.', 'info')
    return redirect(url_for('configuracion') + '#panel-familias')

@app.route('/configuracion/exportar/equipos')
@admin_required
def exportar_equipos_csv():
    flash('Función de exportación en mantenimiento.', 'info')
    return redirect(url_for('configuracion') + '#panel-equipos')

@app.route('/configuracion/exportar/reparaciones')
@admin_required
def exportar_reparaciones_csv():
    flash('Función de exportación en mantenimiento.', 'info')
    return redirect(url_for('configuracion') + '#panel-reparaciones')

@app.route('/health')
def health():
    try:
        db = get_db()
        db.execute('SELECT 1')
        return {'status': 'ok', 'db': 'connected', 'timestamp': datetime.datetime.now().isoformat()}
    except Exception as e:
        return {'status': 'error', 'db': 'disconnected', 'error': str(e)}, 500

@app.route('/backup/crear')
@admin_required
def crear_backup():
    try:
        os.makedirs('backups', exist_ok=True)
        timestamp = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
        backup_path = f'backups/laboratorio_{timestamp}.db'
        shutil.copy2('laboratorio.db', backup_path)
        registrar_historial(session['user_id'], session['username'], 'BACKUP', 'Sistema', None, f'Backup creado: {backup_path}')
        flash(f'Backup creado exitosamente: {backup_path}', 'success')
    except Exception as e:
        flash(f'Error al crear backup: {str(e)}', 'danger')
    return redirect(url_for('configuracion') + '#panel-sistema')

@app.route('/backup/descargar')
@admin_required
def descargar_backup():
    try:
        return send_file('laboratorio.db', as_attachment=True, download_name=f'laboratorio_backup_{datetime.datetime.now().strftime("%Y%m%d")}.db')
    except Exception as e:
        flash(f'Error al descargar: {str(e)}', 'danger')
        return redirect(url_for('configuracion'))

@app.route('/backup/archivo/<filename>')
@admin_required
def descargar_backup_archivo(filename):
    path = os.path.join('backups', filename)
    if not os.path.isfile(path):
        flash('Backup no encontrado.', 'danger')
        return redirect(url_for('configuracion') + '#panel-sistema')
    return send_file(path, as_attachment=True, download_name=filename)

@app.route('/backup/eliminar/<filename>')
@admin_required
def eliminar_backup(filename):
    try:
        path = f'backups/{filename}'
        if os.path.exists(path):
            os.remove(path)
            flash(f'Backup {filename} eliminado', 'success')
            registrar_historial(session['user_id'], session['username'], 'ELIMINAR_BACKUP', 'Sistema', None, f'Eliminado: {filename}')
    except Exception as e:
        flash(f'Error: {str(e)}', 'danger')
    return redirect(url_for('configuracion') + '#panel-sistema')

@app.route('/configuracion/historial')
@admin_required
def ver_historial():
    try:
        db = get_db()

        f_fecha_desde = request.args.get('f_fecha_desde', '')
        f_fecha_hasta = request.args.get('f_fecha_hasta', '')
        f_usuario = request.args.get('f_usuario', '')
        f_accion = request.args.get('f_accion', '')
        f_entidad = request.args.get('f_entidad', '')

        page = max(1, request.args.get('page', 1, type=int))
        per_page = 50

        where_clauses = []
        params = []

        if f_fecha_desde:
            where_clauses.append('fecha >= ?')
            params.append(f_fecha_desde + ' 00:00')
        if f_fecha_hasta:
            where_clauses.append('fecha <= ?')
            params.append(f_fecha_hasta + ' 23:59')
        if f_usuario:
            where_clauses.append('username LIKE ?')
            params.append(f'%{f_usuario}%')
        if f_accion:
            where_clauses.append('accion LIKE ?')
            params.append(f'%{f_accion}%')
        if f_entidad:
            where_clauses.append('entidad LIKE ?')
            params.append(f'%{f_entidad}%')

        where_sql = (' WHERE ' + ' AND '.join(where_clauses)) if where_clauses else ''

        total = db.execute(f'SELECT COUNT(*) as total FROM Historial {where_sql}', params).fetchone()['total']
        total_pages = max(1, (total + per_page - 1) // per_page)
        if page > total_pages: page = total_pages
        offset = (page - 1) * per_page

        historial = db.execute(f'''SELECT * FROM Historial {where_sql}
                                      ORDER BY fecha DESC LIMIT ? OFFSET ?''',
                                      params + [per_page, offset]).fetchall()

        acciones_opciones = ['LOGIN', 'LOGOUT', 'CREAR', 'EDITAR', 'ELIMINAR', 'COMENTARIO', 'CAMBIO_ESTADO', 'ASIGNAR_TECNICO', 'BACKUP', 'CONFIGURACION']
        entidades_opciones = ['Reparacion', 'Usuario', 'Tecnico', 'Catalogo', 'Sistema']

        return render_template('historial.html', historial=historial, page=page, total_pages=total_pages,
                               f_fecha_desde=f_fecha_desde, f_fecha_hasta=f_fecha_hasta,
                               f_usuario=f_usuario, f_accion=f_accion, f_entidad=f_entidad,
                               acciones_opciones=acciones_opciones, entidades_opciones=entidades_opciones)
    except Exception as e:
        import traceback
        flash(f'Error: {str(e)}', 'danger')
        return render_template('error.html', code=500, message=str(e))

@app.route('/admin/limpiar_historial')
@admin_required
def limpiar_historial():
    try:
        db = get_db()
        resultado = db.execute("DELETE FROM Historial WHERE fecha < datetime('now', '-1 year')")
        db.commit()
        count = resultado.rowcount
        registrar_historial(session['user_id'], session['username'], 'LIMPIAR_HISTORIAL', 'Sistema', None, f'Registros eliminados: {count}')
        flash(f'Se eliminaron {count} registros de historial antiguos.', 'success')
    except Exception as e:
        flash(f'Error: {str(e)}', 'danger')
    return redirect(url_for('configuracion') + '#panel-sistema')

@app.route('/exportar/csv')
def exportar_csv():
    db = get_db()

    tabs_estado = list(INDEX_FILTERS)
    is_tecnico = session.get('rol') == 'tecnico'
    if is_tecnico:
        tabs_estado.insert(1, ('MIS_REPARACIONES', 'Mis rep.'))

    estados_validos = {estado for estado, _ in tabs_estado}
    estado_defecto = 'MIS_REPARACIONES' if is_tecnico else 'PEND_REPARACION'
    estado_actual = request.args.get('estado', estado_defecto, type=str).upper()
    if estado_actual not in estados_validos:
        estado_actual = estado_defecto

    busqueda = request.args.get('q', '', type=str).strip()
    f_sala, f_tecnico, f_estado = request.args.get('f_sala', ''), request.args.get('f_tecnico', ''), request.args.get('f_estado', '')
    f_anio, f_mes, f_dia, f_diasemana = request.args.get('f_anio', ''), request.args.get('f_mes', ''), request.args.get('f_dia', ''), request.args.get('f_diasemana', '')

    base_where_clauses, base_query_params = [], []
    if f_sala: base_where_clauses.append('r.sala = ?'); base_query_params.append(f_sala)
    if f_tecnico:
        if f_tecnico == 'SIN_ASIGNAR': base_where_clauses.append('r.tecnico_id IS NULL')
        else: base_where_clauses.append('t.nombre = ?'); base_query_params.append(f_tecnico)
    if f_estado: base_where_clauses.append('r.estado = ?'); base_query_params.append(f_estado)
    if f_anio: base_where_clauses.append('r.fecha LIKE ?'); base_query_params.append(f'%/{f_anio}%')
    if f_mes: base_where_clauses.append('(r.fecha LIKE ? OR r.fecha LIKE ?)'); base_query_params.extend([f'%/{int(f_mes):02d}/%', f'%/{int(f_mes)}/%'])
    if f_dia: base_where_clauses.append('(r.fecha LIKE ? OR r.fecha LIKE ?)'); base_query_params.extend([f'{int(f_dia):02d}/%', f'{int(f_dia)}/%'])
    if f_diasemana: base_where_clauses.append('r.dia_semana = ?'); base_query_params.append(int(f_diasemana))
    if busqueda:
        base_where_clauses.append('(' + ' OR '.join(['r.fecha LIKE ?', 'r.sala LIKE ?', 'r.uid LIKE ?', 'r.npu LIKE ?', 'r.equipo LIKE ?', 'r.estado LIKE ?', 'r.observaciones LIKE ?', 't.nombre LIKE ?']) + ')')
        base_query_params.extend([f'%{busqueda}%'] * 8)

    base_where_sql = (' WHERE ' + ' AND '.join(base_where_clauses)) if base_where_clauses else ''

    tecnico_id_actual = session.get('tecnico_id') or 0

    final_where_clauses, final_query_params = list(base_where_clauses), list(base_query_params)

    if estado_actual == 'URGENTES':
        final_where_clauses.extend(["r.urgente = 'SI'", "UPPER(r.estado) NOT LIKE 'REPARADO%'", "UPPER(r.estado) NOT LIKE 'SIN REPARACION%'"])
    elif estado_actual == 'MIS_REPARACIONES':
        final_where_clauses.extend(["r.tecnico_id = ?", "UPPER(r.estado) NOT LIKE 'REPARADO%'", "UPPER(r.estado) NOT LIKE 'SIN REPARACION%'"])
        final_query_params.append(tecnico_id_actual)
    elif estado_actual == 'PEND_REPARACION':
        final_where_clauses.append('UPPER(r.estado) = ?'); final_query_params.append('PEND. DE REVISION')
    elif estado_actual == 'EN_REPARACION':
        final_where_clauses.append('UPPER(r.estado) = ?'); final_query_params.append('EN REPARACION')
    elif estado_actual == 'REPARADOS':
        final_where_clauses.append('UPPER(r.estado) LIKE ?'); final_query_params.append('REPARADO%')
    elif estado_actual == 'SIN_REPARACION':
        final_where_clauses.append('UPPER(r.estado) LIKE ?'); final_query_params.append('SIN REPARACION%')
    elif estado_actual == 'PENDIENTES':
        final_where_clauses.extend(["UPPER(r.estado) NOT LIKE 'REPARADO%'", "UPPER(r.estado) NOT LIKE 'SIN REPARACION%'", "UPPER(r.estado) != 'PEND. DE REVISION'", "UPPER(r.estado) != 'EN REPARACION'"])

    final_where_sql = (' WHERE ' + ' AND '.join(final_where_clauses)) if final_where_clauses else ''

    query = f'''
        SELECT r.id, r.fecha, r.sala, r.uid, r.npu, r.familia, r.equipo, r.urgente,
               r.estado, r.observaciones, r.fecha_en_reparacion, r.fecha_reparado,
               t.nombre AS tecnico_nombre
        FROM Reparaciones r
        LEFT JOIN Tecnicos t ON r.tecnico_id = t.id
        {final_where_sql}
        ORDER BY r.id DESC
    '''
    reparaciones = db.execute(query, final_query_params).fetchall()

    output = io.StringIO()
    output.write('ID,Fecha,Sala,UID,NPU,Familia,Equipo,Urgente,Tecnico,Estado,Fecha En Rep,Fecha Reparado,Observaciones\n')
    for r in reparaciones:
        obs = (r['observaciones'] or '').replace('\n', ' ').replace(',', ';')
        output.write(f'{r["id"]},{r["fecha"]},{r["sala"]},{r["uid"] or ""},{r["npu"] or ""},{r["familia"] or ""},{r["equipo"]},{r["urgente"]},{r["tecnico_nombre"] or "Sin asignar"},{r["estado"]},{r["fecha_en_reparacion"] or ""},{r["fecha_reparado"] or ""},{obs}\n')

    output.seek(0)
    timestamp = datetime.datetime.now().strftime('%Y%m%d_%H%M')
    return send_file(io.BytesIO(output.getvalue().encode('utf-8')),
                     as_attachment=True,
                     download_name=f'Reporte_{estado_actual}_{timestamp}.csv',
                     mimetype='text/csv')

@app.errorhandler(404)
def not_found(e):
    return render_template('error.html', code=404, message='Página no encontrada'), 404

@app.errorhandler(500)
def server_error(e):
    return render_template('error.html', code=500, message='Error interno del servidor'), 500

@app.errorhandler(400)
def bad_request(e):
    return render_template('error.html', code=400, message='Solicitud inválida'), 400

@app.errorhandler(CSRFError)
def handle_csrf_error(e):
    if request.path.startswith('/api/'):
        return jsonify({'success': False, 'error': 'La sesión de seguridad expiró. Recargá la página e intentá de nuevo.'}), 400
    flash('La sesión de seguridad expiró. Recargá la página e intentá de nuevo.', 'danger')
    return redirect(request.referrer or url_for('index'))

if __name__ == '__main__':
    init_db()

    def backup_scheduler():
        while True:
            time.sleep(86400)
            try:
                os.makedirs('backups', exist_ok=True)
                timestamp = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
                backup_path = f'backups/laboratorio_{timestamp}.db'
                shutil.copy2('laboratorio.db', backup_path)
                print(f'Backup automático creado: {backup_path}')
            except Exception as e:
                print(f'Error en backup automático: {e}')

    threading.Thread(target=backup_scheduler, daemon=True).start()

    app.run(debug=True)
